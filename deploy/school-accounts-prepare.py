#!/usr/bin/env python3
"""Classify the roster, then hash credentials into a new restricted private folder."""
import argparse
import collections
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import unicodedata
import openpyxl


def normalize(value):
    return unicodedata.normalize('NFKC', str(value).strip()).lower() if value is not None else ''


def classify(filename, generate_missing_passwords=False):
    book = openpyxl.load_workbook(filename, data_only=True, read_only=True)
    if len(book.worksheets) != 2:
        raise ValueError('Expected two roster worksheets')
    student_sheet, teacher_sheet = book.worksheets
    expected = ['26-27班別', '班號', '英文姓名 English Name', '中文姓名', '登入名稱', '登入密碼']
    rows = list(student_sheet.iter_rows(values_only=True))
    if [str(v).strip() for v in rows[0][:6]] != expected:
        raise ValueError('Student columns differ from the reviewed layout')
    students, teachers, rejected, generated_students = [], [], [], []
    class_counts = collections.Counter()
    summary_counts = {}
    categories = collections.Counter({'student_header': 1})
    seen_logins, seen_class_numbers = set(), set()
    summary_mode = False
    grand_total = None
    for number, row in enumerate(rows[1:], 2):
        values = list(row[:6]) + [None] * max(0, 6 - len(row))
        cl, class_no, english, chinese, login, password = values
        if not any(v is not None and str(v).strip() for v in values):
            categories['empty'] += 1
            continue
        if str(cl).strip() == '班別' and str(class_no).strip() in ('總人數', '人數') and not any(values[2:]):
            summary_mode = True; categories['summary_header'] += 1; continue
        match = re.fullmatch(r'([1-6])\s*([A-Z])', str(cl or '').strip().upper())
        if summary_mode and not any(values[2:]):
            if match and isinstance(class_no, (int, float)) and int(class_no) == class_no:
                key = match[1] + match[2]
                if key in summary_counts:
                    rejected.append({'row': number, 'reason': 'duplicate_summary_class'})
                summary_counts[key] = int(class_no)
                categories['class_summary'] += 1; continue
            if str(cl).strip() in ('總人數', '總計', '合計', '總數', 'Total', 'TOTAL') and isinstance(class_no, (int, float)):
                grand_total = int(class_no); categories['grand_total'] += 1; continue
        fields_ok = match and isinstance(login, str) and login.strip() and (chinese or english)
        password_ok = isinstance(password, str) and bool(password.strip())
        number_ok = isinstance(class_no, (int, float)) and int(class_no) == class_no and 1 <= class_no <= 99
        if not fields_ok or not number_ok or (not password_ok and not generate_missing_passwords):
            rejected.append({'row': number, 'reason': 'invalid_or_missing_student_fields',
                             'missingColumns': [i + 1 for i, v in enumerate(values) if v is None or not str(v).strip()]})
            continue
        key = normalize(login)
        identity = (int(match[1]), match[2], int(class_no))
        if key in seen_logins or identity in seen_class_numbers:
            rejected.append({'row': number, 'reason': 'duplicate_login_or_class_number'}); continue
        seen_logins.add(key); seen_class_numbers.add(identity)
        if not password_ok:
            generated_students.append({'row': number, 'class': match[1] + match[2], 'source': 'generated_missing_password'})
        students.append({'login': key, 'password': password if password_ok else None,
                         'passwordSource': 'source_workbook' if password_ok else 'generated_missing_password',
                         'sourceRow': number, 'displayName': str(chinese or english).strip(),
                         'grade': int(match[1]), 'cls': match[2], 'classNo': int(class_no)})
        class_counts[match[1] + match[2]] += 1
        categories['student'] += 1
    for number, row in enumerate(teacher_sheet.iter_rows(values_only=True), 1):
        name, login = (list(row[:2]) + [None, None])[:2]
        if name is None and login is None:
            continue
        key = normalize(login)
        if not isinstance(name, str) or not name.strip() or not isinstance(login, str) or not key or key in seen_logins:
            rejected.append({'sheet': 'teacher', 'row': number, 'reason': 'invalid_or_duplicate_teacher'}); continue
        seen_logins.add(key)
        teachers.append({'login': key, 'displayName': name.strip(), 'password': None})
    differences = [{'class': key, 'roster': class_counts[key], 'summary': value}
                   for key, value in summary_counts.items() if class_counts[key] != value]
    missing_summaries = [key for key in class_counts if key not in summary_counts]
    result = {'validStudents': len(students), 'validTeachers': len(teachers), 'classes': dict(sorted(class_counts.items())),
              'rowCategories': dict(categories), 'rejected': rejected, 'summaryTotal': grand_total,
              'classSummaryTotal': sum(summary_counts.values()), 'summaryDifferences': differences,
              'teacherPasswordsMissing': len(teachers), 'generatedMissingStudentPasswords': generated_students,
              'missingSummaryClasses': missing_summaries}
    if grand_total is not None and grand_total != sum(summary_counts.values()):
        result['summaryGrandTotalMismatch'] = True
    book.close()
    return students, teachers, result


def restrict_new_directory(directory):
    if directory.exists() or directory.is_symlink():
        raise ValueError('Use a new dedicated output directory')
    directory.mkdir(mode=0o700)
    if os.name != 'nt':
        return
    query = """$target=$env:MAANSHAN_ACCOUNT_PRIVATE_DIRECTORY
$sids=@((Get-Acl -LiteralPath $target).Access | ForEach-Object {$_.IdentityReference.Translate([System.Security.Principal.SecurityIdentifier]).Value})
@{user=[System.Security.Principal.WindowsIdentity]::GetCurrent().User.Value;sids=$sids}|ConvertTo-Json -Compress
"""
    env = {**os.environ, 'MAANSHAN_ACCOUNT_PRIVATE_DIRECTORY': str(directory)}
    def acl():
        return json.loads(subprocess.check_output(['powershell', '-NoProfile', '-NonInteractive', '-Command', query], env=env, text=True))
    current = acl(); allowed = {current['user'], 'S-1-5-18'}
    subprocess.run(['icacls', str(directory), '/inheritance:r', '/grant:r',
        '*' + current['user'] + ':(OI)(CI)F', '*S-1-5-18:(OI)(CI)F'], check=True, capture_output=True)
    extras = set(acl()['sids']) - allowed
    if extras:
        subprocess.run(['icacls', str(directory), '/remove', *['*' + sid for sid in extras]], check=True, capture_output=True)
    if set(acl()['sids']) != allowed:
        raise ValueError('Private directory ACL verification failed')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', type=Path)
    parser.add_argument('--previous', type=Path)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--generate-missing-passwords', action='store_true')
    parser.add_argument('--accept-summary-mismatch', action='store_true')
    args = parser.parse_args()
    students, teachers, summary = classify(args.input, args.generate_missing_passwords)
    if not args.apply:
        print(json.dumps(summary, ensure_ascii=True)); return
    if summary['rejected'] or (not args.accept_summary_mismatch and
       (summary['summaryDifferences'] or summary.get('summaryGrandTotalMismatch') or summary['missingSummaryClasses'])):
        print(json.dumps(summary, ensure_ascii=True)); raise ValueError('Roster classifications require review')
    if not args.output or not args.output.is_absolute():
        raise ValueError('Private absolute output folder required')
    restrict_new_directory(args.output)
    request = {'students': students, 'teachers': teachers, 'sourceSummary': summary,
       'sourceSha256': hashlib.sha256(args.input.read_bytes()).hexdigest(), 'output': str(args.output),
       'previous': str(args.previous) if args.previous else None}
    result = subprocess.run(['node', str(Path(__file__).with_name('school-accounts-build.cjs'))],
       input=json.dumps(request, ensure_ascii=False).encode('utf-8'), stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
    if result.returncode:
        raise ValueError('Hashing failed; private preparation directory retained for inspection')
    print(result.stdout.decode('utf-8').strip())


if __name__ == '__main__':
    try:
        main()
    except Exception:
        print('{"ok":false,"error":"Account preparation failed; no account data was published"}', file=sys.stderr)
        sys.exit(1)
